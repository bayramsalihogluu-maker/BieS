#!/usr/bin/env python3
"""Turkish tax depreciation & amortisation model for China Finance HQ."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = "0000FF"      # hardcoded input
BLACK = "000000"     # formula
GREEN = "008000"     # link to another sheet
YELLOW = "FFFF00"    # key assumption / fill me in
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")

TRY_FMT = '#,##0;(#,##0);-'
TRY2_FMT = '#,##0.00;(#,##0.00);-'
PCT_FMT = '0.00%'

thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def title(ws, row, text, span=8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=13, bold=True, color="FFFFFF")
    c.fill = HDR_FILL
    for i in range(2, span + 1):
        ws.cell(row=row, column=i).fill = HDR_FILL
    ws.row_dimensions[row].height = 20


def section(ws, row, text, span=8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=11, bold=True)
    c.fill = SUB_FILL
    for i in range(2, span + 1):
        ws.cell(row=row, column=i).fill = SUB_FILL


def note(ws, row, text, span=8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=9, italic=True, color="7F7F7F")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 28


def inp(ws, row, col, value, fmt=None, comment=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=10, color=BLUE, bold=True)
    c.fill = PatternFill("solid", fgColor=YELLOW)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    return c


def frm(ws, row, col, value, fmt=None, color=BLACK):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=10, color=color)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    return c


def lbl(ws, row, col, text, bold=False, size=10):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name=FONT, size=size, bold=bold)
    return c


def hdr_row(ws, row, headers, start=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start + i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = 30


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


wb = Workbook()

# ============================================================ 1. README
ws = wb.active
ws.title = "README"
widths(ws, {"A": 4, "B": 30, "C": 88})
ws.sheet_view.showGridLines = False

r = 1
c = ws.cell(row=r, column=2, value="Türkiye · Tax Depreciation & Amortisation Model")
c.font = Font(name=FONT, size=16, bold=True, color="1F3864")
r += 1
c = ws.cell(row=r, column=2, value="Supporting calculation file for the report to China Finance HQ")
c.font = Font(name=FONT, size=11, italic=True, color="7F7F7F")
r += 2

for label, text in [
    ("Purpose", "Worked examples of depreciation (amortisman) and amortisation of intangibles under the "
                "Turkish Tax Procedure Law (Vergi Usul Kanunu, Law No. 213). Every figure is a live formula: "
                "change a yellow input cell and the schedules recalculate."),
    ("Scope", "Tax book (VUK) treatment only. Statutory/IFRS book figures may differ; the gap is a "
              "temporary difference and is reconciled in the corporate tax return, not in this file."),
    ("Currency", "Turkish Lira (TRY), unless stated otherwise."),
    ("Tax year", "2026 statutory thresholds. Thresholds are re-set every year by the Ministry of Treasury "
                 "and Finance, so this file must be refreshed each January."),
    ("Prepared", "September 2026"),
]:
    lbl(ws, r, 2, label, bold=True)
    c = ws.cell(row=r, column=3, value=text)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30 if len(text) > 95 else 15
    r += 1

r += 1
section(ws, r, "How to read this file", 3)
r += 2

hdr_row(ws, r, ["", "Cell style", "Meaning"], start=1)
r += 1
for style, meaning in [
    ("Yellow fill, blue bold text", "INPUT. Change these. Everything else recalculates."),
    ("Black text", "Formula. Do not overwrite."),
    ("Green text", "Formula linked to the Assumptions tab."),
    ("Grey italic text", "Note, legal reference or source."),
]:
    frm(ws, r, 2, style)
    frm(ws, r, 3, meaning)
    if style.startswith("Yellow"):
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row=r, column=2).font = Font(name=FONT, size=10, color=BLUE, bold=True)
    if style.startswith("Green"):
        ws.cell(row=r, column=2).font = Font(name=FONT, size=10, color=GREEN)
    r += 1

r += 1
section(ws, r, "Tabs", 3)
r += 2
hdr_row(ws, r, ["", "Tab", "What it shows"], start=1)
r += 1
for tab, what in [
    ("Assumptions", "All statutory thresholds and rates for 2026, each with its legal source. Every example reads from here."),
    ("Ex1 Straight-Line", "Normal (straight-line) depreciation of machinery. The default method."),
    ("Ex2 SL vs Declining", "Same asset under straight-line and declining-balance. Shows the timing benefit only."),
    ("Ex3 Passenger Car", "The most error-prone case: 2026 expense caps plus pro-rata first-year depreciation. Two purchase structures compared."),
    ("Ex4 Revaluation", "Optional revaluation under VUK Mük. 298/Ç and its effect on the depreciation base."),
    ("Ex5 Intangibles", "Software, licences, leasehold improvements and goodwill."),
    ("Sources", "Full legal references and links."),
]:
    frm(ws, r, 2, tab)
    frm(ws, r, 3, what)
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 26
    r += 1

r += 1
c = ws.cell(row=r, column=2, value="IMPORTANT")
c.font = Font(name=FONT, size=10, bold=True, color="C00000")
c = ws.cell(row=r, column=3,
            value="This file is a working model, not tax advice. Useful lives are set asset-by-asset and "
                  "sector-by-sector in the official depreciation schedule (VUK General Communiqué No. 333 as amended). "
                  "Confirm the rate for each specific asset against the current schedule before filing.")
c.font = Font(name=FONT, size=10, color="C00000")
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[r].height = 42

# ============================================================ 2. ASSUMPTIONS
ws = wb.create_sheet("Assumptions")
widths(ws, {"A": 42, "B": 18, "C": 12, "D": 66})
ws.sheet_view.showGridLines = False
title(ws, 1, "Assumptions · 2026 statutory parameters", 4)

r = 3
section(ws, r, "A. Thresholds", 4)
r += 1
hdr_row(ws, r, ["Parameter", "Value", "Unit", "Legal source"])
r += 1
A_ROW = {}

rows_a = [
    ("Direct expensing threshold (below this, expense at once)", 12000, "TRY",
     "VUK Art. 313, final para. 2026 amount, VAT excluded. Assets forming one economic/technical unit are tested together.", "threshold"),
    ("Passenger car: depreciable base cap, ÖTV+VAT excluded", 1380000, "TRY",
     "Income Tax Law Art. 40/7; 2026 amounts in Income Tax General Communiqué No. 332.", "car_base_excl"),
    ("Passenger car: depreciable base cap, ÖTV+VAT capitalised", 2600000, "TRY",
     "Income Tax Law Art. 40/7; applies also to second-hand purchases.", "car_base_incl"),
    ("Passenger car: ÖTV+VAT expensed directly, cap", 1200000, "TRY",
     "Income Tax Law Art. 40/1. Alternative to capitalising the taxes.", "car_tax_cap"),
    ("Passenger car: monthly rental cap, VAT excluded", 46000, "TRY",
     "Income Tax Law Art. 40/5. Shown for completeness; not used in the examples.", "car_rent"),
]
for label, val, unit, src, key in rows_a:
    lbl(ws, r, 1, label)
    inp(ws, r, 2, val, TRY_FMT)
    frm(ws, r, 3, unit)
    frm(ws, r, 4, src)
    ws.cell(row=r, column=4).font = Font(name=FONT, size=9, color="7F7F7F")
    ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 26
    A_ROW[key] = r
    r += 1

r += 1
section(ws, r, "B. Methods and limits", 4)
r += 1
hdr_row(ws, r, ["Parameter", "Value", "Unit", "Legal source"])
r += 1
rows_b = [
    ("Declining balance: multiple of normal rate", 2, "x",
     "VUK Repeated Art. 315. Balance-sheet taxpayers only.", "db_mult"),
    ("Declining balance: maximum rate", 0.50, "%",
     "VUK Repeated Art. 315. The doubled rate may never exceed 50%.", "db_cap"),
    ("Useful life: maximum multiple of the official life", 2, "x",
     "VUK Art. 320, as amended by Law No. 7338 (26.10.2021). A longer life may be elected, never a shorter one.", "life_mult"),
    ("Useful life: absolute ceiling", 50, "years",
     "VUK Art. 320. The elected life may not exceed 50 years.", "life_cap"),
    ("Goodwill and formation expenses: amortisation period", 5, "years",
     "VUK Art. 326. Equal instalments.", "gw_life"),
    ("Leasehold improvements: period if lease term unknown", 5, "years",
     "VUK Art. 327. Otherwise over the lease term, in equal percentages. Declining balance is not allowed.", "lh_life"),
]
for label, val, unit, src, key in rows_b:
    lbl(ws, r, 1, label)
    inp(ws, r, 2, val, PCT_FMT if unit == "%" else '#,##0')
    frm(ws, r, 3, unit)
    frm(ws, r, 4, src)
    ws.cell(row=r, column=4).font = Font(name=FONT, size=9, color="7F7F7F")
    ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 26
    A_ROW[key] = r
    r += 1

r += 1
section(ws, r, "C. Indexation", 4)
r += 1
hdr_row(ws, r, ["Parameter", "Value", "Unit", "Legal source"])
r += 1
rows_c = [
    ("Revaluation rate for 2025 (applied in 2026)", 0.2549, "%",
     "VUK General Communiqué No. 585, Official Gazette 27.11.2025. Based on the change in the domestic PPI (Yİ-ÜFE) to October.", "reval_rate"),
    ("Inflation adjustment applied? (1 = yes, 0 = no)", 0, "flag",
     "VUK Provisional Art. 37, added by Law No. 7571 Art. 34, Official Gazette 25.12.2025 No. 33118. "
     "Inflation adjustment is suspended for 2025, 2026 and 2027 regardless of whether the statutory conditions are met.", "infl_flag"),
]
for label, val, unit, src, key in rows_c:
    lbl(ws, r, 1, label)
    inp(ws, r, 2, val, PCT_FMT if unit == "%" else '#,##0')
    frm(ws, r, 3, unit)
    frm(ws, r, 4, src)
    ws.cell(row=r, column=4).font = Font(name=FONT, size=9, color="7F7F7F")
    ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 38
    A_ROW[key] = r
    r += 1

r += 1
section(ws, r, "D. Representative useful lives", 4)
r += 1
hdr_row(ws, r, ["Asset class", "Useful life (years)", "Rate", "Note"])
r += 1
rows_d = [
    ("Buildings (commercial / industrial, reinforced concrete)", 50, "bldg"),
    ("Computers and hardware", 4, "pc"),
    ("Computer software", 3, "sw"),
    ("Intangible rights (patents, know-how, licences)", 15, "ip"),
    ("Passenger cars", 5, "car"),
    ("Office furniture and fixtures", 5, "furn"),
    ("Machinery and equipment (varies widely by sector)", 10, "mach"),
]
for label, life, key in rows_d:
    lbl(ws, r, 1, label)
    inp(ws, r, 2, life, '#,##0')
    frm(ws, r, 3, f"=1/B{r}", PCT_FMT)
    A_ROW[key] = r
    r += 1

note(ws, r, "Useful lives above are the commonly applied figures and are shown as a starting point only. The binding "
            "schedule is the Depreciation List annexed to VUK General Communiqué No. 333 (28.04.2004), as amended by "
            "Communiqués No. 339, 345, 365, 389, 399, 406, 418, 439 and 458. It runs to several hundred lines and is "
            "differentiated by sector: the same machine can carry a different life in two industries. Where an asset "
            "serves more than one sector, the longest life (lowest rate) applies.", 4)

ASSUMP = "Assumptions"


def A(key):
    return f"'{ASSUMP}'!$B${A_ROW[key]}"


# ============================================================ 3. EX1 STRAIGHT LINE
ws = wb.create_sheet("Ex1 Straight-Line")
widths(ws, {"A": 34, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 40})
ws.sheet_view.showGridLines = False
title(ws, 1, "Example 1 · Normal (straight-line) depreciation — machinery", 7)
note(ws, 2, "Normal amortisman. VUK Art. 315. The default method: the same amount every year over the official useful "
            "life. Turkish tax depreciation is computed on full cost; no residual value is deducted.", 7)

r = 4
section(ws, r, "Inputs", 7)
r += 1
lbl(ws, r, 1, "Acquisition cost (VAT excluded)")
inp(ws, r, 2, 4800000, TRY_FMT)
COST1 = f"$B${r}"
r += 1
lbl(ws, r, 1, "Useful life (years)")
inp(ws, r, 2, 10, '#,##0')
LIFE1 = f"$B${r}"
frm(ws, r, 3, f"={A('mach')}", '#,##0', GREEN)
lbl(ws, r, 4, "← official life from Assumptions")
ws.cell(row=r, column=4).font = Font(name=FONT, size=9, italic=True, color="7F7F7F")
r += 1
lbl(ws, r, 1, "Annual rate")
frm(ws, r, 2, f"=1/{LIFE1}", PCT_FMT)
RATE1 = f"$B${r}"
r += 1
lbl(ws, r, 1, "Annual depreciation")
frm(ws, r, 2, f"={COST1}*{RATE1}", TRY_FMT)
r += 2

section(ws, r, "Schedule", 7)
r += 1
hdr_row(ws, r, ["Year", "Opening NBV", "Depreciation", "Accumulated", "Closing NBV", "", "Check"])
hrow = r
r += 1
first = r
for y in range(1, 11):
    frm(ws, r, 1, y, '#,##0')
    if y == 1:
        frm(ws, r, 2, f"={COST1}", TRY_FMT)
    else:
        frm(ws, r, 2, f"=E{r-1}", TRY_FMT)
    frm(ws, r, 3, f"=IF(A{r}>{LIFE1},0,MIN({COST1}*{RATE1},B{r}))", TRY_FMT)
    frm(ws, r, 4, f"=IF(A{r}=1,C{r},D{r-1}+C{r})", TRY_FMT)
    frm(ws, r, 5, f"=B{r}-C{r}", TRY_FMT)
    r += 1
last = r - 1
frm(ws, r, 1, "Total")
ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
frm(ws, r, 3, f"=SUM(C{first}:C{last})", TRY_FMT)
ws.cell(row=r, column=3).font = Font(name=FONT, size=10, bold=True)
frm(ws, r, 7, f'=IF(ROUND(C{r}-{COST1},2)=0,"OK: fully depreciated","CHECK")')
r += 2
note(ws, r, "Election available since Law No. 7338: the taxpayer may choose a longer life than the official one, up to "
            "twice that life and never beyond 50 years, applying the same rate each year (VUK Art. 320). A shorter life "
            "can never be elected. Since 2021 depreciation may also be computed on a daily basis (VUK Art. 320/A) "
            "instead of taking a full year in the year of acquisition.", 7)

# ============================================================ 4. EX2 SL vs DB
ws = wb.create_sheet("Ex2 SL vs Declining")
widths(ws, {"A": 10, "B": 16, "C": 16, "D": 16, "E": 6, "F": 16, "G": 16, "H": 16, "I": 6, "J": 18})
ws.sheet_view.showGridLines = False
title(ws, 1, "Example 2 · Straight-line vs declining balance — same asset", 10)
note(ws, 2, "Azalan bakiyeler usulü. VUK Repeated Art. 315. The rate is twice the normal rate but capped at 50%, applied "
            "to the net book value each year. The balance remaining in the final year is written off in full. Available "
            "to balance-sheet taxpayers only. A taxpayer may switch from declining balance to straight-line, but never "
            "the other way.", 10)

r = 4
section(ws, r, "Inputs", 10)
r += 1
lbl(ws, r, 1, "Cost")
inp(ws, r, 2, 4800000, TRY_FMT)
COST2 = f"$B${r}"
r += 1
lbl(ws, r, 1, "Useful life")
inp(ws, r, 2, 10, '#,##0')
LIFE2 = f"$B${r}"
r += 1
lbl(ws, r, 1, "Normal rate")
frm(ws, r, 2, f"=1/{LIFE2}", PCT_FMT)
SLRATE = f"$B${r}"
r += 1
lbl(ws, r, 1, "Declining rate")
frm(ws, r, 2, f"=MIN({SLRATE}*{A('db_mult')},{A('db_cap')})", PCT_FMT, GREEN)
DBRATE = f"$B${r}"
lbl(ws, r, 3, "min(2 × normal rate, 50%)")
ws.cell(row=r, column=3).font = Font(name=FONT, size=9, italic=True, color="7F7F7F")
r += 2

section(ws, r, "Comparison", 10)
r += 1
hdr_row(ws, r, ["Year", "SL depreciation", "SL accumulated", "SL closing NBV", "",
                "DB opening NBV", "DB depreciation", "DB closing NBV", "", "Difference (DB − SL)"])
r += 1
f2 = r
for y in range(1, 11):
    frm(ws, r, 1, y, '#,##0')
    frm(ws, r, 2, f"=IF(A{r}>{LIFE2},0,{COST2}*{SLRATE})", TRY_FMT)
    frm(ws, r, 3, f"=IF(A{r}=1,B{r},C{r-1}+B{r})", TRY_FMT)
    frm(ws, r, 4, f"={COST2}-C{r}", TRY_FMT)
    if y == 1:
        frm(ws, r, 6, f"={COST2}", TRY_FMT)
    else:
        frm(ws, r, 6, f"=H{r-1}", TRY_FMT)
    # final year: write off the whole remaining balance
    frm(ws, r, 7, f"=IF(A{r}={LIFE2},F{r},F{r}*{DBRATE})", TRY_FMT)
    frm(ws, r, 8, f"=F{r}-G{r}", TRY_FMT)
    frm(ws, r, 10, f"=G{r}-B{r}", TRY_FMT)
    r += 1
l2 = r - 1
frm(ws, r, 1, "Total")
ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
frm(ws, r, 2, f"=SUM(B{f2}:B{l2})", TRY_FMT)
frm(ws, r, 7, f"=SUM(G{f2}:G{l2})", TRY_FMT)
frm(ws, r, 10, f"=SUM(J{f2}:J{l2})", TRY_FMT)
for col in (2, 7, 10):
    ws.cell(row=r, column=col).font = Font(name=FONT, size=10, bold=True)
r += 1
frm(ws, r, 1, "Check")
frm(ws, r, 10, f'=IF(ROUND(J{r-1},2)=0,"OK: same total, different timing","CHECK")')
ws.cell(row=r, column=10).font = Font(name=FONT, size=10, bold=True, color="006100")
r += 2
note(ws, r, "Both methods deduct the same total cost. Declining balance moves deductions forward, which is a cash-flow "
            "and net-present-value benefit, not a permanent tax saving. In a high-inflation and high-interest environment "
            "that timing benefit is economically significant.", 10)

# ============================================================ 5. EX3 PASSENGER CAR
ws = wb.create_sheet("Ex3 Passenger Car")
widths(ws, {"A": 46, "B": 18, "C": 18, "D": 4, "E": 54})
ws.sheet_view.showGridLines = False
title(ws, 1, "Example 3 · Passenger car — expense caps and pro-rata first year", 5)
note(ws, 2, "Passenger cars carry two restrictions that apply to no other asset: a cap on the deductible cost "
            "(Income Tax Law Art. 40/7 and 40/1) and pro-rata depreciation in the year of acquisition (VUK Art. 320). "
            "Depreciation on the portion above the cap is permanently non-deductible (kanunen kabul edilmeyen gider), "
            "so it is a permanent difference, not a timing difference.", 5)

r = 4
section(ws, r, "Inputs", 5)
r += 1
lbl(ws, r, 1, "List price, ÖTV and VAT excluded")
inp(ws, r, 2, 2000000, TRY_FMT)
CAR_NET = f"$B${r}"
r += 1
lbl(ws, r, 1, "ÖTV + VAT on the vehicle")
inp(ws, r, 2, 1600000, TRY_FMT)
CAR_TAX = f"$B${r}"
r += 1
lbl(ws, r, 1, "Month of acquisition (1 = January)")
inp(ws, r, 2, 4, '#,##0')
CAR_M = f"$B${r}"
r += 1
lbl(ws, r, 1, "Useful life (years)")
inp(ws, r, 2, 5, '#,##0')
CAR_LIFE = f"$B${r}"
r += 1
lbl(ws, r, 1, "Months depreciated in year 1")
frm(ws, r, 2, f"=13-{CAR_M}", '#,##0')
CAR_MONTHS = f"$B${r}"
frm(ws, r, 5, "Pro-rata rule: depreciation starts in the month of acquisition.")
ws.cell(row=r, column=5).font = Font(name=FONT, size=9, italic=True, color="7F7F7F")
r += 2

section(ws, r, "Option A · ÖTV and VAT expensed directly", 5)
r += 1
hdr_row(ws, r, ["Item", "Amount", "Deductible", "", "Basis"])
r += 1
lbl(ws, r, 1, "ÖTV + VAT expensed in year 1")
frm(ws, r, 2, f"={CAR_TAX}", TRY_FMT)
frm(ws, r, 3, f"=MIN({CAR_TAX},{A('car_tax_cap')})", TRY_FMT, GREEN)
frm(ws, r, 5, "Cap: Income Tax Law Art. 40/1")
ws.cell(row=r, column=5).font = Font(name=FONT, size=9, color="7F7F7F")
A_TAX = f"$C${r}"
r += 1
lbl(ws, r, 1, "Depreciable cost")
frm(ws, r, 2, f"={CAR_NET}", TRY_FMT)
frm(ws, r, 3, f"=MIN({CAR_NET},{A('car_base_excl')})", TRY_FMT, GREEN)
frm(ws, r, 5, "Cap: Income Tax Law Art. 40/7, ÖTV+VAT excluded")
ws.cell(row=r, column=5).font = Font(name=FONT, size=9, color="7F7F7F")
A_BASE = f"$C${r}"
r += 1
lbl(ws, r, 1, "Disallowed cost (permanent difference)")
frm(ws, r, 2, f"=MAX(0,{CAR_NET}-{A_BASE})+MAX(0,{CAR_TAX}-{A_TAX})", TRY_FMT)
ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True, color="C00000")
r += 2

section(ws, r, "Option A · depreciation schedule", 5)
r += 1
hdr_row(ws, r, ["Year", "Deductible depreciation", "Note", "", ""])
r += 1
fa = r
for y in range(1, 7):
    frm(ws, r, 1, y, '#,##0')
    if y == 1:
        frm(ws, r, 2, f"={A_BASE}/{CAR_LIFE}*{CAR_MONTHS}/12", TRY_FMT)
        frm(ws, r, 3, "Pro-rata: part year only")
    elif y == 6:
        frm(ws, r, 2, f"={A_BASE}-SUM(B{fa}:B{r-1})", TRY_FMT)
        frm(ws, r, 3, "Balance not taken in year 1 is deducted in the final year")
    else:
        frm(ws, r, 2, f"=IF(A{r}>{CAR_LIFE},0,{A_BASE}/{CAR_LIFE})", TRY_FMT)
        frm(ws, r, 3, "Full year")
    ws.cell(row=r, column=3).font = Font(name=FONT, size=9, italic=True, color="7F7F7F")
    r += 1
la = r - 1
frm(ws, r, 1, "Total")
ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
frm(ws, r, 2, f"=SUM(B{fa}:B{la})", TRY_FMT)
ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True)
frm(ws, r, 3, f'=IF(ROUND(B{r}-{A_BASE},2)=0,"OK: equals capped base","CHECK")')
OPT_A_TOTAL = f"$B${r}"
OPT_A_Y1 = f"$B${fa}"
r += 2

section(ws, r, "Option B · ÖTV and VAT capitalised into cost", 5)
r += 1
hdr_row(ws, r, ["Item", "Amount", "Deductible", "", "Basis"])
r += 1
lbl(ws, r, 1, "Total capitalised cost")
frm(ws, r, 2, f"={CAR_NET}+{CAR_TAX}", TRY_FMT)
frm(ws, r, 3, f"=MIN({CAR_NET}+{CAR_TAX},{A('car_base_incl')})", TRY_FMT, GREEN)
frm(ws, r, 5, "Cap: Income Tax Law Art. 40/7, taxes included")
ws.cell(row=r, column=5).font = Font(name=FONT, size=9, color="7F7F7F")
B_BASE = f"$C${r}"
r += 1
lbl(ws, r, 1, "Disallowed cost (permanent difference)")
frm(ws, r, 2, f"=MAX(0,{CAR_NET}+{CAR_TAX}-{B_BASE})", TRY_FMT)
ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True, color="C00000")
r += 2

section(ws, r, "Option B · depreciation schedule", 5)
r += 1
hdr_row(ws, r, ["Year", "Deductible depreciation", "Note", "", ""])
r += 1
fb = r
for y in range(1, 7):
    frm(ws, r, 1, y, '#,##0')
    if y == 1:
        frm(ws, r, 2, f"={B_BASE}/{CAR_LIFE}*{CAR_MONTHS}/12", TRY_FMT)
        frm(ws, r, 3, "Pro-rata: part year only")
    elif y == 6:
        frm(ws, r, 2, f"={B_BASE}-SUM(B{fb}:B{r-1})", TRY_FMT)
        frm(ws, r, 3, "Balance deducted in the final year")
    else:
        frm(ws, r, 2, f"=IF(A{r}>{CAR_LIFE},0,{B_BASE}/{CAR_LIFE})", TRY_FMT)
        frm(ws, r, 3, "Full year")
    ws.cell(row=r, column=3).font = Font(name=FONT, size=9, italic=True, color="7F7F7F")
    r += 1
lb = r - 1
frm(ws, r, 1, "Total")
ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
frm(ws, r, 2, f"=SUM(B{fb}:B{lb})", TRY_FMT)
ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True)
frm(ws, r, 3, f'=IF(ROUND(B{r}-{B_BASE},2)=0,"OK: equals capped base","CHECK")')
OPT_B_TOTAL = f"$B${r}"
r += 2

section(ws, r, "Which option deducts more?", 5)
r += 1
lbl(ws, r, 1, "Option A: total deduction over life (incl. taxes expensed in year 1)")
frm(ws, r, 2, f"={OPT_A_TOTAL}+{A_TAX}", TRY_FMT)
r += 1
lbl(ws, r, 1, "Option B: total deduction over life")
frm(ws, r, 2, f"={OPT_B_TOTAL}", TRY_FMT)
r += 1
lbl(ws, r, 1, "Advantage of A over B", bold=True)
frm(ws, r, 2, f"=B{r-2}-B{r-1}", TRY_FMT)
ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True)
r += 1
note(ws, r, "The choice is made once, at acquisition, and cannot be revisited. Option A also front-loads the deduction "
            "because the taxes are expensed immediately rather than spread over five years. Run both before buying: with "
            "the 2026 caps the answer changes with the price point.", 5)

# ============================================================ 6. EX4 REVALUATION
ws = wb.create_sheet("Ex4 Revaluation")
widths(ws, {"A": 46, "B": 20, "C": 20, "D": 4, "E": 54})
ws.sheet_view.showGridLines = False
title(ws, 1, "Example 4 · Optional revaluation under VUK Repeated Art. 298/Ç", 5)
note(ws, 2, "Inflation adjustment is suspended for 2025, 2026 and 2027 (VUK Provisional Art. 37, added by Law No. 7571). "
            "In its place, taxpayers may elect to revalue depreciable assets each year by the official revaluation rate. "
            "The uplift is credited to a special fund in equity and is NOT taxed; depreciation is then computed on the "
            "revalued amount. This is the main mechanism protecting Turkish depreciation deductions from inflation.", 5)

r = 4
section(ws, r, "Inputs", 5)
r += 1
lbl(ws, r, 1, "Original cost")
inp(ws, r, 2, 4800000, TRY_FMT)
RV_COST = f"$B${r}"
r += 1
lbl(ws, r, 1, "Accumulated depreciation before revaluation")
inp(ws, r, 2, 1440000, TRY_FMT)
RV_ACC = f"$B${r}"
r += 1
lbl(ws, r, 1, "Useful life (years)")
inp(ws, r, 2, 10, '#,##0')
RV_LIFE = f"$B${r}"
r += 1
lbl(ws, r, 1, "Revaluation rate")
frm(ws, r, 2, f"={A('reval_rate')}", PCT_FMT, GREEN)
RV_RATE = f"$B${r}"
frm(ws, r, 5, "2025 rate: 25.49% (VUK General Communiqué No. 585)")
ws.cell(row=r, column=5).font = Font(name=FONT, size=9, color="7F7F7F")
r += 2

section(ws, r, "Effect of the election", 5)
r += 1
hdr_row(ws, r, ["Item", "Before", "After revaluation", "", "Note"])
r += 1
lbl(ws, r, 1, "Gross cost")
frm(ws, r, 2, f"={RV_COST}", TRY_FMT)
frm(ws, r, 3, f"={RV_COST}*(1+{RV_RATE})", TRY_FMT)
frm(ws, r, 5, "Both cost and accumulated depreciation are indexed")
ws.cell(row=r, column=5).font = Font(name=FONT, size=9, color="7F7F7F")
NEW_COST = f"$C${r}"
r += 1
lbl(ws, r, 1, "Accumulated depreciation")
frm(ws, r, 2, f"={RV_ACC}", TRY_FMT)
frm(ws, r, 3, f"={RV_ACC}*(1+{RV_RATE})", TRY_FMT)
r += 1
lbl(ws, r, 1, "Net book value")
frm(ws, r, 2, f"=B{r-2}-B{r-1}", TRY_FMT)
frm(ws, r, 3, f"=C{r-2}-C{r-1}", TRY_FMT)
r += 1
lbl(ws, r, 1, "Increase credited to special fund (untaxed)", bold=True)
frm(ws, r, 3, f"=C{r-1}-B{r-1}", TRY_FMT)
ws.cell(row=r, column=3).font = Font(name=FONT, size=10, bold=True, color="006100")
frm(ws, r, 5, "Shown in equity. Taxed only if distributed or withdrawn from the business.")
ws.cell(row=r, column=5).font = Font(name=FONT, size=9, color="7F7F7F")
ws.cell(row=r, column=5).alignment = Alignment(wrap_text=True)
r += 1
lbl(ws, r, 1, "Annual depreciation from now on", bold=True)
frm(ws, r, 2, f"={RV_COST}/{RV_LIFE}", TRY_FMT)
frm(ws, r, 3, f"={NEW_COST}/{RV_LIFE}", TRY_FMT)
for col in (2, 3):
    ws.cell(row=r, column=col).font = Font(name=FONT, size=10, bold=True)
r += 1
lbl(ws, r, 1, "Additional annual deduction")
frm(ws, r, 2, f"=C{r-1}-B{r-1}", TRY_FMT)
ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True, color="006100")
r += 2
note(ws, r, "Why this matters to a parent company: without revaluation, depreciation stays fixed in nominal lira while "
            "replacement cost rises with inflation, so the real value of the deduction erodes every year and taxable "
            "profit is overstated. The election is optional and must be assessed annually.", 5)

# ============================================================ 7. EX5 INTANGIBLES
ws = wb.create_sheet("Ex5 Intangibles")
widths(ws, {"A": 40, "B": 16, "C": 14, "D": 16, "E": 60})
ws.sheet_view.showGridLines = False
title(ws, 1, "Example 5 · Intangibles, leasehold improvements and goodwill", 5)
note(ws, 2, "Classification drives the outcome here more than any calculation does. The same payment can be written off "
            "over 3 years or over 15 depending on whether it is treated as software or as a licence.", 5)

r = 4
hdr_row(ws, r, ["Item", "Cost", "Period (years)", "Annual charge", "Legal basis and comment"])
r += 1
rows_e = [
    ("Computer software", 900000, f"={A('sw')}",
     "Depreciation List item 4.3: 3 years. Purchased software, not a licence."),
    ("Licence / intangible right", 900000, f"={A('ip')}",
     "Patents, know-how, licences and similar rights: 15 years. Same cash outflow, five times the write-off period."),
    ("Leasehold improvements (lease term known)", 2400000, 8,
     "VUK Art. 327: over the lease term, in equal percentages. Declining balance is not permitted."),
    ("Leasehold improvements (lease term unknown)", 2400000, f"={A('lh_life')}",
     "VUK Art. 327: 5 years when the term is not fixed."),
    ("Goodwill (peştemallık)", 3000000, f"={A('gw_life')}",
     "VUK Art. 326: 5 years, equal instalments."),
    ("Formation and organisation expenses", 600000, f"={A('gw_life')}",
     "VUK Art. 326: 5 years, equal instalments. May alternatively be expensed as incurred."),
]
first_e = r
for label, cost, period, basis in rows_e:
    lbl(ws, r, 1, label)
    inp(ws, r, 2, cost, TRY_FMT)
    if isinstance(period, str):
        frm(ws, r, 3, period, '#,##0', GREEN)
    else:
        inp(ws, r, 3, period, '#,##0')
    frm(ws, r, 4, f"=B{r}/C{r}", TRY_FMT)
    frm(ws, r, 5, basis)
    ws.cell(row=r, column=5).font = Font(name=FONT, size=9, color="7F7F7F")
    ws.cell(row=r, column=5).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    r += 1

r += 1
section(ws, r, "Software vs licence — the cost of misclassification", 5)
r += 1
lbl(ws, r, 1, "Annual charge if treated as software")
frm(ws, r, 2, f"=D{first_e}", TRY_FMT)
r += 1
lbl(ws, r, 1, "Annual charge if treated as a licence")
frm(ws, r, 2, f"=D{first_e+1}", TRY_FMT)
r += 1
lbl(ws, r, 1, "Annual difference", bold=True)
frm(ws, r, 2, f"=B{r-2}-B{r-1}", TRY_FMT)
ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True, color="C00000")
r += 1
note(ws, r, "Both write off the same total cost, so this is a timing difference. In a high-inflation economy, however, "
            "deductions deferred by twelve years are worth a fraction of their nominal value. Classification should be "
            "documented at the point of purchase, in the contract wording, not decided later by the accounting team.", 5)

# ============================================================ 8. SOURCES
ws = wb.create_sheet("Sources")
widths(ws, {"A": 6, "B": 44, "C": 84})
ws.sheet_view.showGridLines = False
title(ws, 1, "Legal sources", 3)
r = 3
hdr_row(ws, r, ["#", "Instrument", "Relevance"])
r += 1
srcs = [
    ("Tax Procedure Law No. 213 (VUK), Arts. 313–321", "Definition of depreciable assets, the direct expensing threshold, normal depreciation, depreciation period and the pro-rata rule for passenger cars."),
    ("VUK Repeated Art. 315", "Declining balance method; twice the normal rate, capped at 50%."),
    ("VUK Art. 320, as amended by Law No. 7338 (26.10.2021)", "Election of a longer useful life, up to twice the official life and never beyond 50 years."),
    ("VUK Art. 320/A, added by Law No. 7338", "Option to compute depreciation on a daily basis."),
    ("VUK Art. 326", "Goodwill and formation expenses: 5 years, equal instalments."),
    ("VUK Art. 327", "Leasehold improvements: over the lease term in equal percentages; 5 years if the term is unknown."),
    ("VUK Repeated Art. 298/Ç", "Optional continuous revaluation of depreciable assets; the uplift is not taxed."),
    ("VUK Provisional Art. 37, added by Law No. 7571 Art. 34", "Inflation adjustment suspended for 2025, 2026 and 2027. Official Gazette 25.12.2025, No. 33118. Extendable by Presidential decision for up to three further periods."),
    ("VUK General Communiqué No. 333 (28.04.2004), as amended by Nos. 339, 345, 365, 389, 399, 406, 418, 439, 458", "The binding Depreciation List: useful lives and rates, asset by asset and sector by sector."),
    ("VUK General Communiqué No. 585 (Official Gazette 27.11.2025)", "2025 revaluation rate: 25.49%."),
    ("Income Tax Law No. 193, Art. 40/1, 40/5 and 40/7", "Passenger car expense restrictions: direct expense cap, rental cap and depreciable cost cap."),
    ("Income Tax General Communiqué No. 332", "2026 passenger car threshold amounts."),
]
for i, (inst, rel) in enumerate(srcs, 1):
    frm(ws, r, 1, i, '#,##0')
    frm(ws, r, 2, inst)
    frm(ws, r, 3, rel)
    for col in (2, 3):
        ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 32
    r += 1

r += 1
note(ws, r, "Primary legislation is published in the Official Gazette (Resmî Gazete) and consolidated by the Revenue "
            "Administration (Gelir İdaresi Başkanlığı, gib.gov.tr). Threshold amounts are re-set annually; the figures in "
            "this file are the 2026 amounts and must be refreshed each January.", 3)

wb.save("/home/user/BieS/raporlar/TR-Depreciation-Model-2026.xlsx")
print("written")
